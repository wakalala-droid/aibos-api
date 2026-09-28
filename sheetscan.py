"""
AIBOS — Sheet scanner: read EVERY sheet of a workbook the way a person reads it.

The old reader asked pandas for one sheet and assumed row 1 was the header. Real
books from real businesses are almost never shaped that way, and a village-banking
workbook proved it: 15 sheets, headers on rows 2/3/4/24/33/39, several tables on
one sheet, months running ACROSS the page instead of down it, navigation sheets
named "Input Tabs>>>", and a hidden sheet. AI-BOS picked one sheet, read row 1 as
the header, and handed the owner 26 columns called "Unnamed: 0" … "Unnamed: 25".
The other fourteen sheets were thrown away without a word.

This module never assumes. For each sheet it:
  • fills merged cells, so a header spanning four columns names all four,
  • splits the sheet into TABLES (runs of rows; one blank row is tolerated inside
    a table because a header is so often separated from its data by exactly one),
  • finds each table's header ROW by what the rows around it look like, never by
    position,
  • recognises a MATRIX (Name | Jul | Aug | Sep …) and turns it back into one row
    per name per month, which is the only shape a transaction can be booked in,
  • marks total rows, so a "Total" line is never imported as a transaction,
  • counts what is actually filled in, so a blank template says it is blank
    instead of importing ten thousand zeroes.

Pure and dependency-light (openpyxl/pandas, both already required), so every bit
of it is testable without a database, a network, or an AI key.
"""

import io
import re
import logging
from datetime import datetime, date, time

log = logging.getLogger("aibos.sheetscan")

# Guard rails. A corrupt or hostile file must not be able to spend the server's
# memory: past these the scan stops and says it truncated.
MAX_ROWS_PER_SHEET = 50000
MAX_COLS_PER_SHEET = 256
MAX_TABLES_PER_SHEET = 40

# How many rows at the top of a block may turn out to be the header.
HEADER_LOOKAHEAD = 8
# A table ends at this many consecutive blank rows. One blank row is tolerated,
# because a header is so often separated from its first data row by exactly one.
BLANK_ROWS_END_TABLE = 2
# A run of this many date-like headers across the page means a matrix layout.
MIN_PERIOD_RUN = 3

_TOTAL_WORDS = (
    "total", "totals", "grand total", "sub total", "subtotal", "sum",
    "net", "balance", "closing", "opening", "average", "avg", "cumulative",
)
# Headers of columns holding a figure worked out FROM the other columns. On a
# matrix they are kept out of the unpivot, so a year's total is never imported
# alongside the twelve months that add up to it.
_DERIVED_WORDS = ("total", "sum", "cumulative", "balance", "average", "avg", "ytd", "grand")

_MONTHS = ("jan", "feb", "mar", "apr", "may", "jun",
           "jul", "aug", "sep", "oct", "nov", "dec")
_MONTH_YEAR = re.compile(
    r"^\s*(" + "|".join(_MONTHS) + r")[a-z]*\.?[\s\-/]*((?:19|20)?\d{2})?\s*$", re.I)
_YEAR_MONTH = re.compile(r"^\s*(?:19|20)\d{2}[-/](?:0?[1-9]|1[0-2])(?:[-/]\d{1,2})?\s*$")
_PERIOD_WORD = re.compile(r"^\s*(?:q[1-4]|quarter\s*[1-4]|week\s*\d+|wk\s*\d+|p\d{1,2})\b", re.I)
# A sheet that only exists to divide the workbook up: "Input Tabs>>>".
_DIVIDER = re.compile(r"[<>]{2,}|^-{3,}$|^={3,}$")


# ══════════════════════════════════════════════════════════════════════════════
# Cell predicates
# ══════════════════════════════════════════════════════════════════════════════

def is_blank(v) -> bool:
    if v is None:
        return True
    if isinstance(v, str):
        return v.strip() == ""
    # NaN is the only value not equal to itself; pandas leaves them everywhere.
    return v != v


def is_number(v) -> bool:
    if isinstance(v, bool) or is_blank(v):
        return False
    if isinstance(v, (int, float)):
        return v == v
    if isinstance(v, (datetime, date, time)):
        return False
    s = str(v).strip().replace(",", "").replace("%", "")
    s = re.sub(r"^[A-Za-z$£€]{0,3}\s*", "", s)   # K 1,500 / ZMW 1500 / $12
    s = re.sub(r"^\((.*)\)$", r"-\1", s)                   # (1,500) is negative
    try:
        float(s)
        return True
    except ValueError:
        return False


def is_date_like(v) -> bool:
    """A real date cell, or text a person would read as a period ("Jul-26", "Q3")."""
    if isinstance(v, bool):
        return False
    if isinstance(v, (datetime, date)):
        return True
    if is_blank(v) or isinstance(v, (int, float)):
        return False
    s = str(v).strip()
    return bool(_MONTH_YEAR.match(s) or _YEAR_MONTH.match(s) or _PERIOD_WORD.match(s))


def text_of(v) -> str:
    if is_blank(v):
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat() if (v.hour, v.minute, v.second) == (0, 0, 0) else v.isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return str(v).strip()


def _norm(s) -> str:
    return re.sub(r"[^a-z0-9 ]+", " ", text_of(s).lower()).strip()


def looks_like_total(label) -> bool:
    """A 'Total'/'Balance' line, worked out from other lines, which must never be
    imported as a transaction of its own."""
    n = _norm(label)
    if not n:
        return False
    return any(n == w or n.startswith(w + " ") or n.endswith(" " + w) for w in _TOTAL_WORDS)


def _is_derived_header(label) -> bool:
    n = _norm(label)
    return bool(n) and any(w in n for w in _DERIVED_WORDS)


def col_letter(i: int) -> str:
    """0 → A, 25 → Z, 26 → AA. Named after what the owner sees in Excel."""
    s = ""
    i += 1
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


def to_number(v):
    """A cell as a number, reading the way money is actually written down:
    "K 1,500", "(1,500)" for negative, "12%"."""
    if isinstance(v, bool) or is_blank(v):
        return None
    if isinstance(v, (int, float)):
        return float(v) if v == v else None
    if isinstance(v, (datetime, date, time)):
        return None
    s = text_of(v).replace(",", "")
    neg = bool(re.match(r"^\(.*\)$", s))
    s = re.sub(r"^\((.*)\)$", r"\1", s)
    s = re.sub(r"^[A-Za-z$£€]{0,3}\s*", "", s).replace("%", "").strip()
    try:
        n = float(s)
    except ValueError:
        return None
    return -n if neg else n


def period_iso(v) -> str | None:
    """A period header as an ISO date: a real date cell, or text like "Jul-26"."""
    if isinstance(v, bool):
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = text_of(v)
    m = _MONTH_YEAR.match(s)
    if m:
        month = _MONTHS.index(m.group(1)[:3].lower()) + 1
        year = m.group(2)
        if year:
            year = int(year)
            year += 2000 if year < 100 else 0
        else:
            year = datetime.now().year
        return f"{year:04d}-{month:02d}-01"
    if _YEAR_MONTH.match(s):
        parts = re.split(r"[-/]", s.strip())
        y, mo = int(parts[0]), int(parts[1])
        d = int(parts[2]) if len(parts) > 2 else 1
        try:
            return date(y, mo, d).isoformat()
        except ValueError:
            return None
    return None


# ══════════════════════════════════════════════════════════════════════════════
# Loading — one grid per sheet, merged cells filled, padding trimmed
# ══════════════════════════════════════════════════════════════════════════════

def _fill_merges(ws, grid: list) -> None:
    """A merged block reads as its top-left value at EVERY cell it covers.

    openpyxl gives the value once and None everywhere else, so a header merged
    across four columns looked like one header and three blank columns, and the
    three columns under it lost their names."""
    for rng in list(ws.merged_cells.ranges):
        r0, r1 = rng.min_row - 1, rng.max_row - 1
        c0, c1 = rng.min_col - 1, rng.max_col - 1
        if r0 < 0 or c0 < 0 or r0 >= len(grid) or c0 >= len(grid[r0]):
            continue
        v = grid[r0][c0]
        if is_blank(v):
            continue
        for r in range(r0, min(r1 + 1, len(grid))):
            for c in range(c0, min(c1 + 1, len(grid[r]))):
                if is_blank(grid[r][c]):
                    grid[r][c] = v


def _trim(grid: list) -> list:
    """Drop the empty rows and columns Excel pads a sheet out to (often 1000×26)."""
    last_row = -1
    last_col = -1
    for r, row in enumerate(grid):
        for c, v in enumerate(row):
            if not is_blank(v):
                last_row = r
                last_col = max(last_col, c)
    if last_row < 0:
        return []
    width = last_col + 1
    out = []
    for row in grid[:last_row + 1]:
        cut = list(row[:width])
        cut += [None] * (width - len(cut))
        out.append(cut)
    return out


def load_grids(content: bytes, filename: str = "") -> tuple:
    """(ordered {sheet: grid}, {sheet: meta}) for any workbook or delimited file."""
    ext = (filename or "").rsplit(".", 1)[-1].lower()

    if ext in ("csv", "txt", "tsv"):
        return _load_delimited(content)

    if ext == "xls":
        return _load_via_pandas(content, engine="xlrd")

    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    except Exception as exc:  # noqa: BLE001 — .xlsb, .ods, a mislabelled file
        log.info("sheetscan: openpyxl could not open %s (%s); trying pandas", filename, exc)
        return _load_via_pandas(content)

    grids, meta = {}, {}
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(
            max_row=min(ws.max_row or 0, MAX_ROWS_PER_SHEET),
            max_col=min(ws.max_column or 0, MAX_COLS_PER_SHEET),
            values_only=True)]
        _fill_merges(ws, rows)
        grids[ws.title] = _trim(rows)
        meta[ws.title] = {
            "hidden": ws.sheet_state != "visible",
            "images": len(getattr(ws, "_images", []) or []),
            "truncated": bool((ws.max_row or 0) > MAX_ROWS_PER_SHEET
                              or (ws.max_column or 0) > MAX_COLS_PER_SHEET),
        }
    try:
        wb.close()
    except Exception:  # noqa: BLE001
        pass
    return grids, meta


def _load_via_pandas(content: bytes, engine: str = "openpyxl") -> tuple:
    import pandas as pd
    xl = pd.ExcelFile(io.BytesIO(content), engine=engine)
    grids, meta = {}, {}
    for name in xl.sheet_names:
        # header=None: the header is found later, from the shape of the rows.
        df = xl.parse(name, header=None, nrows=MAX_ROWS_PER_SHEET)
        grid = [[None if (v != v) else v for v in row] for row in df.values.tolist()]
        grids[name] = _trim(grid)
        meta[name] = {"hidden": False, "images": 0, "truncated": False}
    return grids, meta


def _sniff_separator(text: str) -> str | None:
    """Which character separates the fields, judged over the WHOLE file.

    pandas sniffs from the first line. A CSV whose first line is a title
    ("Shop report") has no commas on it, so the sniffer split on spaces and the
    columns came back as "Shop" and "report". The right separator is the one
    that gives the same field count on the most lines."""
    lines = [ln for ln in text.splitlines()[:60] if ln.strip()]
    if not lines:
        return None
    best, best_score = None, 0
    for sep in (",", ";", "	", "|"):
        counts = [ln.count(sep) for ln in lines if ln.count(sep) > 0]
        if len(counts) < 2:
            continue
        modal = max(set(counts), key=counts.count)
        # How many lines agree on that field count, and how many fields it makes.
        score = counts.count(modal) * (modal + 1)
        if modal >= 1 and score > best_score:
            best, best_score = sep, score
    return best


def _load_delimited(content: bytes) -> tuple:
    import pandas as pd
    last = None
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            text = content.decode(enc)
        except (UnicodeDecodeError, LookupError) as e:
            last = e
            continue
        sep = _sniff_separator(text)
        # A title line above the header has fewer fields than the table under
        # it, and pandas takes the FIRST line as the width: every later row then
        # "saw 3 fields, expected 1" and the whole file failed to open. Naming
        # the widest row's worth of columns up front lets the ragged top of a
        # real-world CSV through, where the header is the fourth line down.
        names = None
        if sep:
            import csv as _csv
            sample = list(_csv.reader(io.StringIO(text[:400000]), delimiter=sep))
            width = max((len(r) for r in sample[:500]), default=1)
            names = list(range(width))
        try:
            df = pd.read_csv(io.StringIO(text),
                             sep=sep if sep else None,
                             engine="python", header=None, names=names, dtype=object,
                             nrows=MAX_ROWS_PER_SHEET)
        except Exception as e:  # noqa: BLE001
            last = e
            continue
        grid = [[None if (v != v) else v for v in row] for row in df.values.tolist()]
        return ({"Sheet1": _trim(grid)},
                {"Sheet1": {"hidden": False, "images": 0, "truncated": False}})
    raise ValueError(f"Could not read that file as CSV ({last}). Try saving it as CSV UTF-8.")


# ══════════════════════════════════════════════════════════════════════════════
# Blocks → tables
# ══════════════════════════════════════════════════════════════════════════════

def _blocks(grid: list) -> list:
    """(start, end) row ranges, split where the sheet goes properly quiet."""
    filled = [any(not is_blank(v) for v in row) for row in grid]
    out, start, gap = [], None, 0
    for i, has in enumerate(filled):
        if has:
            if start is None:
                start = i
            gap = 0
        elif start is not None:
            gap += 1
            if gap >= BLANK_ROWS_END_TABLE:
                out.append((start, i - gap))
                start, gap = None, 0
    if start is not None:
        out.append((start, len(grid) - 1 - gap))
    return [(a, b) for a, b in out if b >= a]


def _row_stats(row: list) -> dict:
    filled = [v for v in row if not is_blank(v)]
    return {
        "filled": len(filled),
        "labelish": sum(1 for v in filled if not is_number(v)),
        "numeric": sum(1 for v in filled if is_number(v)),
        "dates": sum(1 for v in filled if is_date_like(v)),
    }


def _pick_header(grid: list, start: int, end: int) -> int:
    """Which row of this block names the columns.

    Scored on what a header actually is: mostly words or dates rather than
    figures, wide enough to cover the data, and followed by rows that are more
    numeric than it is. Position is only the tie-breaker, which is why a header
    on row 4 under a title on row 3 is found at all."""
    best, best_score = start, float("-inf")
    widest = max((_row_stats(grid[r])["filled"] for r in range(start, end + 1)), default=0)
    for r in range(start, min(start + HEADER_LOOKAHEAD, end + 1)):
        st = _row_stats(grid[r])
        if st["filled"] < 2:
            continue                       # a one-cell line is a title, not a header
        below = [_row_stats(grid[x]) for x in range(r + 1, min(r + 6, end + 1))]
        if not below:
            continue
        below_filled = sum(b["filled"] for b in below) / len(below)
        below_num = sum(b["numeric"] for b in below) / max(1, sum(b["filled"] for b in below))
        score = (
            2.0 * (st["labelish"] + st["dates"]) / st["filled"]     # words/dates, not figures
            + 1.0 * min(st["filled"], widest) / max(1, widest)       # as wide as the data
            + 2.0 * below_num                                        # figures underneath it
            + (0.7 if below_filled >= st["filled"] * 0.5 else 0.0)
            - 0.05 * (r - start)                                     # earliest wins a tie
        )
        if score > best_score:
            best, best_score = r, score
    return best


def _title_above(grid: list, header_row: int, block_start: int) -> str:
    """The nearest short line of text above the header — the table's own name."""
    for r in range(header_row - 1, max(block_start - 3, -1), -1):
        st = _row_stats(grid[r])
        if 0 < st["filled"] <= 2 and st["labelish"]:
            t = re.sub(r"\s+", " ", text_of(next((v for v in grid[r] if not is_blank(v)), ""))).strip()
            if t and not _DIVIDER.search(t) and not is_number(t):
                return t[:80]
    return ""


def _used_columns(grid: list, header: list, data_rows: list) -> list:
    """The columns this table actually occupies.

    A sheet is as wide as its widest row, so a narrow table sitting on a wide
    sheet was given every column to its right — twenty "Column Q" … "Column AD"
    that hold nothing. Only columns with something in the header or in one of
    this table's own rows count."""
    width = max(len(header), max((len(grid[r]) for r in data_rows), default=0))
    used = []
    for c in range(width):
        if (c < len(header) and not is_blank(header[c])) or \
           any(c < len(grid[r]) and not is_blank(grid[r][c]) for r in data_rows):
            used.append(c)
    return used


def _column_names(header: list) -> list:
    """Header cells → unique, readable column names; blanks named by their letter."""
    names, seen = [], {}
    for i, v in enumerate(header):
        name = re.sub(r"\s+", " ", text_of(v)).strip() or f"Column {col_letter(i)}"
        key = name.lower()
        if key in seen:
            seen[key] += 1
            name = f"{name} ({seen[key]})"
        else:
            seen[key] = 1
        names.append(name)
    return names


def _period_run(header: list) -> tuple:
    """The longest stretch of date-like headers across the page → (start, end),
    or (0, -1) when the sheet is not laid out as a matrix."""
    best = (0, -1)
    run_start = None
    for i in range(len(header) + 1):
        if i < len(header) and is_date_like(header[i]):
            if run_start is None:
                run_start = i
        elif run_start is not None:
            if (i - 1) - run_start > best[1] - best[0]:
                best = (run_start, i - 1)
            run_start = None
    return best if best[1] - best[0] + 1 >= MIN_PERIOD_RUN else (0, -1)


# ══════════════════════════════════════════════════════════════════════════════
# One table
# ══════════════════════════════════════════════════════════════════════════════


# ── A cash book: one header up top, money in two columns, dates written once ──

# Header names that mean "money came in" and "money went out". Matched whole,
# not by substring: a column called "Invoice" must never read as "in".
_MONEY_IN_HEADERS = (
    "in", "money in", "cash in", "in s", "ins", "received", "receipts", "receipt",
    "credit", "credits", "deposit", "deposits", "inflow", "inflows", "paid in",
)
_MONEY_OUT_HEADERS = (
    "out", "money out", "cash out", "out s", "outs", "paid", "paid out", "payment",
    "payments", "spent", "debit", "debits", "withdrawal", "withdrawals",
    "outflow", "outflows", "expense", "expenses",
)
# A running balance is never the amount of a transaction.
_BALANCE_HEADERS = ("balance", "bal", "running balance", "closing balance", "b/f", "c/f")


def is_balance_header(name) -> bool:
    return _norm(name) in _BALANCE_HEADERS


def _money_pair(names: list, used: list) -> tuple | None:
    """(in column, out column) when the table keeps money in two columns."""
    ins = [c for c in used if _norm(names[c]) in _MONEY_IN_HEADERS]
    outs = [c for c in used if _norm(names[c]) in _MONEY_OUT_HEADERS]
    if len(ins) == 1 and len(outs) == 1:
        return ins[0], outs[0]
    return None


def _date_column(grid: list, data_rows: list, used: list) -> int | None:
    """The column holding the date, judged by what is in it."""
    for c in used:
        vals = [grid[r][c] for r in data_rows if c < len(grid[r]) and not is_blank(grid[r][c])]
        if vals and sum(1 for v in vals if is_date_like(v)) >= max(1, len(vals) * 0.5):
            return c
    return None


def _carried_dates(grid: list, data_rows: list, col: int) -> dict:
    """{row: the date that row belongs to}, blanks inheriting the date above.

    A ledger writes the date once and the next few lines belong to it. Read
    literally, those lines have no date and land nowhere in the timeline."""
    out, last = {}, None
    for r in data_rows:
        v = grid[r][col] if col < len(grid[r]) else None
        if not is_blank(v) and is_date_like(v):
            last = v
        out[r] = last
    return out


def looks_like_header_row(row: list) -> bool:
    """True when this row NAMES the columns, rather than being one of them.

    A header is words. The moment a candidate carries a date or more than one
    figure it is a line of the ledger, and promoting it to a header costs the
    real column names and one row of real data."""
    st = _row_stats(row)
    if st["filled"] < 2 or st["dates"]:
        return False
    return st["labelish"] / st["filled"] >= 0.7 and st["numeric"] <= 1

def _split_in_out(grid: list, names: list, data_rows: list, used: list,
                  in_col: int, out_col: int, date_col: int | None, dates: dict) -> list:
    """One row per money figure, tagged with which way the money went.

    A cash-book line can be both ("sale 15,000 in, 15,000 banked out"), and both
    are real, so it becomes two entries rather than one of them being dropped."""
    out = []
    for r in data_rows:
        row = grid[r]
        base = {names[c]: (row[c] if c < len(row) else None) for c in used}
        if date_col is not None:
            base[names[date_col]] = dates.get(r)
        first_label = next((text_of(row[c]) for c in used
                            if c < len(row) and not is_blank(row[c]) and not is_number(row[c])), "")
        is_total = looks_like_total(first_label)
        for col, direction in ((in_col, "in"), (out_col, "out")):
            v = row[col] if col < len(row) else None
            n = to_number(v)
            if n is None or n == 0:
                continue
            out.append({**base, "Direction": direction, "Amount": abs(n),
                        "_row": r + 1, "_is_total": is_total})
    return out


def _mostly_numeric(grid: list, data_rows: list, col: int) -> bool:
    vals = [grid[r][col] for r in data_rows if col < len(grid[r]) and not is_blank(grid[r][col])]
    if not vals:
        return False
    return sum(1 for v in vals if is_number(v)) > len(vals) * 0.6


def _row_has_value(row: dict) -> bool:
    if "Amount" in row:
        return bool(to_number(row.get("Amount")))
    return any(to_number(v) for k, v in row.items() if not k.startswith("_"))


def _unpivot(grid: list, header: list, names: list, data_rows: list,
             p0: int, p1: int, used: list) -> tuple:
    """Name | Jul | Aug | Sep → one row per name per month.

    A matrix is how people lay a year out on a page, and the only shape AI-BOS
    cannot book, because a transaction needs ONE date and ONE amount."""
    label_cols = [c for c in used if c < p0
                  and any(c < len(grid[r]) and not is_blank(grid[r][c]) for r in data_rows)]
    after = [c for c in used if c > p1]
    dropped = [names[c] for c in after if _is_derived_header(names[c])]
    # A second block of months to the right of this one belongs to its own table,
    # not repeated onto every row of this one.
    dropped += [names[c] for c in after
                if is_date_like(header[c]) and names[c] not in dropped]
    keep_after = [c for c in after if names[c] not in dropped]
    # The date each period column stands for, read once from the header row.
    periods = {c: (period_iso(header[c]) or names[c]) for c in range(p0, p1 + 1)}

    out = []
    for r in data_rows:
        row = grid[r]
        labels = {names[c]: (text_of(row[c]) if c < len(row) else "") for c in label_cols}
        extra = {names[c]: (row[c] if c < len(row) else None) for c in keep_after}
        is_total = looks_like_total(next((v for v in labels.values() if v), ""))
        for c in range(p0, p1 + 1):
            v = row[c] if c < len(row) else None
            if is_blank(v):
                continue
            n = to_number(v)
            out.append({
                **labels, **extra,
                "Date": periods[c],
                "Amount": n if n is not None else text_of(v),
                "_row": r + 1,
                "_col": col_letter(c),
                "_is_total": is_total,
            })
    columns = [names[c] for c in label_cols] + [names[c] for c in keep_after] + ["Date", "Amount"]
    return out, columns, [names[c] for c in label_cols], dropped


def _flat_rows(grid: list, names: list, data_rows: list, used: list,
               date_col: int | None = None, dates: dict | None = None) -> list:
    out = []
    for r in data_rows:
        row = grid[r]
        rec = {names[c]: (row[c] if c < len(row) else None) for c in used}
        if date_col is not None and dates:
            rec[names[date_col]] = dates.get(r)
        first_label = next((text_of(row[c]) for c in used
                            if c < len(row) and not is_blank(row[c]) and not is_number(row[c])), "")
        rec["_row"] = r + 1
        rec["_is_total"] = looks_like_total(first_label)
        out.append(rec)
    return out


def _build_table(sheet: str, grid: list, start: int, end: int, index: int,
                 inherit: dict | None = None) -> dict | None:
    if inherit:
        # A continuation of the table above: it has no header of its own,
        # and its first line is data, not column names.
        header_row = inherit["header_row"] - 1
        header, names, used = inherit["header"], inherit["names"], inherit["used"]
        data_rows = [r for r in range(start, end + 1)
                     if any(not is_blank(v) for v in grid[r])]
        if not data_rows:
            return None
        p0, p1 = (0, -1)
        matrix = False
    else:
        header_row = _pick_header(grid, start, end)
        header = grid[header_row]
        data_rows = [r for r in range(header_row + 1, end + 1)
                     if any(not is_blank(v) for v in grid[r])]
        if not data_rows:
            return None

        used = _used_columns(grid, header, data_rows)
        if not used:
            return None
        header = list(header[:used[-1] + 1]) + [None] * max(0, used[-1] + 1 - len(header))
        names = _column_names(header)
        p0, p1 = _period_run([header[c] if c in set(used) else None
                              for c in range(len(header))])
        matrix = p1 >= p0

    table = {
        "id": f"{sheet}!{header_row + 1}",
        "sheet": sheet,
        "title": _title_above(grid, header_row, start) or sheet,
        "index": index,
        "header_row": header_row + 1,            # 1-based, as Excel shows it
        "first_data_row": data_rows[0] + 1,
        "last_data_row": data_rows[-1] + 1,
        "orientation": "matrix" if matrix else "rows",
        "dropped_columns": [],
        "notes": [],
    }

    if matrix:
        rows, columns, label_names, dropped = _unpivot(grid, header, names, data_rows, p0, p1, used)
        table["columns"] = columns
        table["label_columns"] = label_names
        table["period_count"] = p1 - p0 + 1
        table["dropped_columns"] = dropped
        if dropped:
            table["notes"].append(
                "Columns worked out from the months (" + ", ".join(dropped[:4]) +
                ") were left out, so a total is not imported on top of the months it adds up.")
    else:
        date_col = _date_column(grid, data_rows, used)
        dates = _carried_dates(grid, data_rows, date_col) if date_col is not None else {}
        pair = _money_pair(names, used)
        if pair:
            rows = _split_in_out(grid, names, data_rows, used,
                                 pair[0], pair[1], date_col, dates)
            table["columns"] = [names[c] for c in used] + ["Direction", "Amount"]
            table["orientation"] = "cashbook"
            table["notes"].append(
                "Money is kept in two columns here (" + names[pair[0]] + " and "
                + names[pair[1]] + "), so each figure is read as its own entry and "
                "the running balance is left alone.")
        else:
            rows = _flat_rows(grid, names, data_rows, used, date_col, dates)
        table["columns"] = table.get("columns") or [names[c] for c in used]
        table["label_columns"] = [names[c] for c in used
                                  if not _mostly_numeric(grid, data_rows, c)]

    table["rows"] = rows
    table["row_count"] = len(rows)
    table["total_rows"] = sum(1 for r in rows if r.get("_is_total"))
    table["nonzero_rows"] = sum(1 for r in rows if _row_has_value(r))
    table["all_zero"] = table["row_count"] > 0 and table["nonzero_rows"] == 0
    if table["all_zero"]:
        table["notes"].append(
            "Every figure in this table is blank or zero — it looks like an unused template.")
    # Kept for a continuation block to inherit; stripped before the table travels.
    table["_header"], table["_names"], table["_used"] = header, names, used
    return table


# ══════════════════════════════════════════════════════════════════════════════
# The whole workbook
# ══════════════════════════════════════════════════════════════════════════════

def _sheet_is_divider(grid: list, name: str) -> str | None:
    """Why this sheet holds nothing to import, or None when it does."""
    if not grid:
        return "The sheet is empty."
    cells = [v for row in grid for v in row if not is_blank(v)]
    if not cells:
        return "The sheet is empty."
    if len(cells) <= 3 and any(_DIVIDER.search(text_of(v)) for v in cells + [name]):
        return "A divider sheet that only labels a section of the workbook."
    if len(cells) < 3:
        return "The sheet has almost nothing in it."
    return None



def _shares_columns(grid: list, start: int, end: int, used: list) -> bool:
    """True when this block sits in the same columns as the table above it."""
    here = set()
    for r in range(start, end + 1):
        for c, v in enumerate(grid[r]):
            if not is_blank(v):
                here.add(c)
    if not here:
        return False
    return len(here & set(used)) / len(here) >= 0.6


def scan(content: bytes, filename: str = "", sample_rows: int = 0) -> dict:
    """Every sheet, every table, ready to map.

    `sample_rows` trims each table's rows for a preview that has to travel back
    over the web; 0 keeps them all, which is what the import itself uses."""
    grids, meta = load_grids(content, filename)
    sheets, tables = [], []

    for name, grid in grids.items():
        skip = _sheet_is_divider(grid, name)
        info = {
            "name": name,
            "hidden": meta.get(name, {}).get("hidden", False),
            "images": meta.get(name, {}).get("images", 0),
            "truncated": meta.get(name, {}).get("truncated", False),
            "skipped": bool(skip),
            "reason": skip or "",
            "tables": [],
        }
        if not skip:
            last = None          # the last table on this sheet with a real header
            for i, (a, b) in enumerate(_blocks(grid)[:MAX_TABLES_PER_SHEET]):
                # A section under a header that was written once, at the top of the
                # sheet, and separated from the next by a blank line or a running
                # total. Its first line is data; reading it as a header costs the
                # real column names and a row of real figures.
                carry = (last is not None
                         and not looks_like_header_row(grid[_pick_header(grid, a, b)])
                         and _shares_columns(grid, a, b, last["used"]))
                try:
                    t = _build_table(name, grid, a, b, i, last if carry else None)
                except Exception as exc:  # noqa: BLE001 — one odd table never stops the rest
                    log.warning("sheetscan: %s rows %s-%s failed: %s", name, a, b, exc)
                    continue
                if not t or not t["rows"]:
                    continue
                if carry and tables and tables[-1]["id"] == last["id"]:
                    # One ledger, not five: the sections belong together.
                    prev = tables[-1]
                    prev["rows"] = prev["rows"] + t["rows"]
                    prev["row_count"] = len(prev["rows"])
                    prev["total_rows"] = sum(1 for r in prev["rows"] if r.get("_is_total"))
                    prev["nonzero_rows"] = sum(1 for r in prev["rows"] if _row_has_value(r))
                    prev["all_zero"] = prev["row_count"] > 0 and prev["nonzero_rows"] == 0
                    prev["last_data_row"] = t["last_data_row"]
                    continue
                if not carry:
                    last = {"id": t["id"], "header_row": t["header_row"],
                            "header": t["_header"], "names": t["_names"],
                            "used": t["_used"]}
                info["tables"].append(t["id"])
                if sample_rows and len(t["rows"]) > sample_rows:
                    t = {**t, "rows": t["rows"][:sample_rows], "sampled": True}
                tables.append(t)
            if not info["tables"]:
                info["skipped"] = True
                info["reason"] = "No table of figures was found on this sheet."
        sheets.append(info)

    return {
        "sheets": sheets,
        "tables": tables,
        "sheet_count": len(sheets),
        "table_count": len(tables),
        "row_total": sum(t["row_count"] for t in tables),
        "nonzero_total": sum(t["nonzero_rows"] for t in tables),
    }


def find_table(scanned: dict, table_id: str) -> dict | None:
    return next((t for t in scanned.get("tables", []) if t["id"] == table_id), None)
