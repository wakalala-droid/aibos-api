"""
AIBOS — The AIBOS workbook: a spreadsheet laid out the way AIBOS reads.

Any spreadsheet can be uploaded, and AIBOS works hard to understand one it has
never seen (sheetscan.py, doc_ai.py, attach.py). But a new owner with no
spreadsheet yet should not have to invent one, and a guess is still a guess.
So AIBOS hands them its own: one sheet for each kind of money, headed the way
the importer files it, with dropdowns for how it was paid and what it was for.

  • build()        the workbook, made for this business: their products and
                   workers in the dropdowns, their own payment methods and
                   biggest costs first (from the interview, interview.py).
  • recognise()    a table that came back from it, known by its headings, so
                   it is filed exactly: no AI call and no guessing.
  • resolve()      every line of a recognised sheet as the event it is. The
                   owner chose the sheet, the date, the amount and the
                   category, so a line nothing had to be worked out for is
                   booked like an entry typed on the Record page. A wage to a
                   name not on the worker list still asks, as any import does.
  • row_key()      what makes one line that line, so the same workbook can be
                   kept, added to and uploaded again: lines AIBOS already has
                   are left out, and only the new ones go in.

Recognition is by headings, not by a hidden mark, so a sheet copied into
another workbook or through Google Sheets is still known.
"""

import hashlib
import io
import logging
import re

import attach
import sheetscan

log = logging.getLogger("aibos.starter_book")

TAG = "AIBOS workbook v1"
FILENAME = "AIBOS workbook.xlsx"

# Lines a workbook keeps room for, styled and offered dropdowns.
ROWS = 2000

# Columns: (heading, width, kind). Kinds style the column and pick a dropdown.
SHEETS = (
    {"kind": "sales", "title": "Sales", "event_type": "Sale",
     "what": "Your sales from the AIBOS workbook",
     "columns": (("Date", 13, "date"), ("Customer", 24, "text"), ("What was sold", 30, "product"),
                 ("Quantity", 11, "qty"), ("Amount", 14, "money"), ("Paid by", 16, "paid_by"),
                 ("Notes", 32, "text")),
     "required": ("date", "customer", "what was sold", "amount"),
     "mapping": {"date": "Date", "amount": "Amount", "counterparty": "Customer",
                 "description": "What was sold", "quantity": "Quantity"},
     "guide": "Sales: what you sold, to whom and for how much. A day's takings on one line is fine."},
    {"kind": "expenses", "title": "Expenses", "event_type": "Expense",
     "what": "Your running costs from the AIBOS workbook",
     "columns": (("Date", 13, "date"), ("Paid to", 24, "text"), ("What it was for", 30, "text"),
                 ("Category", 22, "category"), ("Amount", 14, "money"), ("Paid by", 16, "paid_by"),
                 ("Notes", 32, "text")),
     "required": ("date", "paid to", "what it was for", "category", "amount"),
     "mapping": {"date": "Date", "amount": "Amount", "counterparty": "Paid to",
                 "description": "What it was for", "category": "Category"},
     "guide": "Expenses: rent, fuel, airtime, repairs and every other running cost."},
    {"kind": "stock", "title": "Stock bought", "event_type": "Purchase",
     "what": "Stock you bought, from the AIBOS workbook",
     "columns": (("Date", 13, "date"), ("Supplier", 24, "text"), ("Item", 30, "product"),
                 ("Quantity", 11, "qty"), ("Amount", 14, "money"), ("Paid by", 16, "paid_by"),
                 ("Notes", 32, "text")),
     "required": ("date", "supplier", "item", "quantity", "amount"),
     "mapping": {"date": "Date", "amount": "Amount", "counterparty": "Supplier",
                 "description": "Item", "quantity": "Quantity"},
     "guide": "Stock bought: goods you buy to sell, or to use in what you make. Amount is the total you paid."},
    {"kind": "wages", "title": "Wages", "event_type": "Salary",
     "what": "Wages from the AIBOS workbook",
     "columns": (("Date", 13, "date"), ("Worker", 26, "worker"), ("Amount", 14, "money"),
                 ("Paid by", 16, "paid_by"), ("Notes", 32, "text")),
     "required": ("date", "worker", "amount"),
     "mapping": {"date": "Date", "amount": "Amount", "counterparty": "Worker"},
     "guide": "Wages: what you paid each worker, one line per payment."},
    {"kind": "products", "title": "Products", "event_type": None,
     "what": "Your products, for your product list",
     "columns": (("Product", 30, "text"), ("Category", 18, "text"), ("Unit", 12, "unit"),
                 ("Buying price", 14, "money"), ("Selling price", 14, "money"),
                 ("In stock now", 13, "qty"), ("Reorder at", 12, "qty")),
     "required": ("product", "selling price"),
     "mapping": {"name": "Product", "category": "Category", "unit": "Unit",
                 "buy_price": "Buying price", "sell_price": "Selling price",
                 "opening_stock": "In stock now", "reorder_level": "Reorder at"},
     "guide": "Products: what you sell, with prices and how many you have today."},
    {"kind": "workers", "title": "Workers", "event_type": None,
     "what": "The people you pay, for your worker list",
     "columns": (("Name", 28, "text"), ("Job", 22, "text"), ("Monthly pay", 14, "money"),
                 ("Phone", 16, "text"), ("Started on", 13, "date")),
     "required": ("name", "job", "monthly pay"),
     "mapping": {"name": "Name", "position": "Job", "basic_pay": "Monthly pay",
                 "phone": "Phone", "start_date": "Started on"},
     "guide": "Workers: the people you pay. AIBOS works out their PAYE and NAPSA from here."},
)
BY_KIND = {s["kind"]: s for s in SHEETS}
RECORD_KINDS = ("products", "workers")

PAID_BY = ("Cash", "Mobile money", "Bank transfer", "Card", "Cheque", "On credit")

# What the owner picks under Category → what AIBOS books. A category not here
# is kept in the owner's own words.
CATEGORIES = (
    ("Rent", "Expense", "rent"),
    ("Fuel and transport", "Expense", "transport"),
    ("Electricity and water", "Expense", "utilities"),
    ("Airtime and data", "Expense", "communication"),
    ("Repairs and maintenance", "Expense", "repairs"),
    ("Loan repayments", "Loan", None),
    ("Tax and levies", "TaxPayment", None),
    ("Cleaning", "Expense", "cleaning"),
    ("Security", "Expense", "security"),
    ("Marketing", "Expense", "marketing"),
    ("Bank charges", "Expense", "bank charges"),
    ("Professional fees", "Expense", "professional"),
    ("Insurance", "Expense", "insurance"),
    ("Licences and permits", "Expense", "licences"),
    ("Food and refreshments", "Expense", "food"),
    ("Packaging", "Expense", "packaging"),
    ("Stationery and printing", "Expense", "stationery"),
    ("Training", "Expense", "training"),
    ("Subscriptions", "Expense", "subscriptions"),
    ("Other", "Expense", "general"),
)
_CATEGORY = {c[0].lower(): c for c in CATEGORIES}

# Each industry's own costs, offered near the top.
INDUSTRY_CATEGORIES = {
    "Hospitality": (("Laundry", "Expense", "laundry"), ("Guest supplies", "Expense", "housekeeping")),
    "Restaurant / Food": (("Gas and charcoal", "Expense", "utilities"),),
    "Agriculture": (("Seed, fertiliser and feed", "Expense", "farm inputs"), ("Vet and medicines", "Expense", "farm inputs")),
    "Transport": (("Tyres and parts", "Expense", "repairs"),),
    "Manufacturing": (("Raw materials", "Expense", "raw materials"),),
}

# The interview's "biggest costs" → the Category they mean.
_BIG_COST_CATEGORY = {
    "Rent": "Rent", "Fuel and transport": "Fuel and transport",
    "Electricity and water": "Electricity and water", "Airtime and data": "Airtime and data",
    "Loan repayments": "Loan repayments", "Repairs": "Repairs and maintenance",
}
_PAID_BY_INTERVIEW = {"Cash": "Cash", "Mobile money": "Mobile money",
                      "Bank transfer": "Bank transfer", "Card": "Card"}

UNITS = ("each", "kg", "litre", "bag", "box", "crate", "carton", "dozen", "pack",
         "bottle", "metre", "night", "hour", "plate")


def _norm(s) -> str:
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9 ]+", " ", str(s or "").lower())).strip()


def _first(values, preferred) -> list:
    """`values` with the `preferred` ones first, in the order preferred gives them."""
    lead = [v for v in preferred if v in values]
    return lead + [v for v in values if v not in lead]


def categories_for(industry: str | None, answers: dict | None = None) -> list:
    """The Category dropdown: the owner's biggest costs first, then their
    industry's own, then the rest."""
    rows = list(INDUSTRY_CATEGORIES.get(industry or "", ())) + [c for c in CATEGORIES]
    names = []
    for r in rows:
        if r[0] not in names:
            names.append(r[0])
    big = ((answers or {}).get("big_costs") or {}).get("value") or []
    lead = [_BIG_COST_CATEGORY[b] for b in big if b in _BIG_COST_CATEGORY]
    lead += [r[0] for r in INDUSTRY_CATEGORIES.get(industry or "", ())]
    return _first(names, lead)


def paid_by_for(answers: dict | None = None) -> list:
    said = ((answers or {}).get("paid_by") or {}).get("value") or []
    return _first(list(PAID_BY), [_PAID_BY_INTERVIEW[p] for p in said if p in _PAID_BY_INTERVIEW])


# ══════════════════════════════════════════════════════════════════════════════
# Making the workbook
# ══════════════════════════════════════════════════════════════════════════════

def build(business: str | None = None, industry: str | None = None, answers: dict | None = None,
          products: list | None = None, workers: list | None = None) -> bytes:
    """The workbook as .xlsx bytes, made for this business."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    ink = "1C1917"
    head_font = Font(name="Calibri", bold=True, color=ink, size=11)
    head_fill = PatternFill("solid", fgColor="EEF2F7")
    rule = Border(bottom=Side(style="thin", color="94A3B8"))
    name = (business or "").strip() or "Your business"

    wb = Workbook()
    wb.properties.title = f"{name} AIBOS workbook"
    wb.properties.creator = "AIBOS"
    wb.properties.keywords = TAG

    # ── Start here ──────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Start here"
    ws.column_dimensions["A"].width = 100
    ws.sheet_view.showGridLines = False
    lines = [
        (f"{name}: AIBOS workbook", Font(bold=True, size=16, color=ink)),
        ("Fill this in, then upload it on the Upload a file page. AIBOS files every "
         "line in the right place, with no questions about columns.", Font(size=11, color="44403C")),
        ("", None),
        ("How to fill it in", Font(bold=True, size=12, color=ink)),
        ("1. One line per sale, cost or payment.", None),
        ("2. Dates as day/month/year, for example 14/10/2026.", None),
        ("3. Amounts as plain numbers: 2500, not K2,500.", None),
        ("4. Paid by and Category offer a list. Pick from it, or type your own.", None),
        ("5. Leave a sheet empty if it does not apply to you. Keep the sheet names and the "
         "column headings as they are.", None),
        ("6. Keep adding to this same file and upload it again whenever you like: AIBOS "
         "records only the lines it does not have yet.", None),
        ("", None),
        ("The sheets", Font(bold=True, size=12, color=ink)),
    ] + [(s["guide"], None) for s in SHEETS]
    for i, (text, font) in enumerate(lines, start=1):
        cell = ws.cell(row=i, column=1, value=text or None)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if font:
            cell.font = font

    # ── The lists behind the dropdowns (hidden) ─────────────────────────────
    lists = wb.create_sheet("Lists")
    cats = categories_for(industry, answers)
    paid = paid_by_for(answers)
    for col, (heading, values) in enumerate((("Paid by", paid), ("Category", cats), ("Unit", list(UNITS))),
                                            start=1):
        lists.cell(row=1, column=col, value=heading).font = head_font
        for r, v in enumerate(values, start=2):
            lists.cell(row=r, column=col, value=v)
    lists.sheet_state = "hidden"
    list_ref = {
        "paid_by": f"Lists!$A$2:$A${len(paid) + 1}",
        "category": f"Lists!$B$2:$B${len(cats) + 1}",
        "unit": f"Lists!$C$2:$C${len(UNITS) + 1}",
        "product": f"Products!$A$2:$A${ROWS}",
        "worker": f"Workers!$A$2:$A${ROWS}",
    }

    # ── One sheet for each kind of money ────────────────────────────────────
    sheets = {}
    for spec in SHEETS:
        sh = wb.create_sheet(spec["title"])
        sheets[spec["kind"]] = sh
        sh.freeze_panes = "A2"
        for c, (heading, width, kind) in enumerate(spec["columns"], start=1):
            cell = sh.cell(row=1, column=c, value=heading)
            cell.font, cell.fill, cell.border = head_font, head_fill, rule
            cell.alignment = Alignment(vertical="center")
            letter = cell.column_letter
            dim = sh.column_dimensions[letter]
            dim.width = width
            # Styled by column, not by cell: the sheet stays as short as the
            # owner's own lines, and Excel formats each new one as it is typed.
            if kind == "date":
                dim.number_format = "DD/MM/YYYY"
            elif kind == "money":
                dim.number_format = "#,##0.00"
            elif kind == "qty":
                dim.number_format = "#,##0.##"
            if kind in list_ref:
                # A dropdown that offers, never refuses: their own words are fine.
                # No leading "=": the file format keeps formulas without one.
                dv = DataValidation(type="list", formula1=list_ref[kind],
                                    allow_blank=True, showErrorMessage=False)
                sh.add_data_validation(dv)
                dv.add(f"{letter}2:{letter}{ROWS}")
        sh.row_dimensions[1].height = 22

    # What the business already has, so the dropdowns name them exactly.
    for r, p in enumerate((products or [])[:ROWS - 1], start=2):
        row = (p.get("name"), p.get("category"), p.get("unit"), p.get("buy_price"),
               p.get("sell_price"), None, p.get("reorder_level"))
        for c, v in enumerate(row, start=1):
            if v not in (None, "", 0):
                sheets["products"].cell(row=r, column=c, value=v)
    for r, e in enumerate((workers or [])[:ROWS - 1], start=2):
        row = (e.get("name"), e.get("position"), e.get("basic_pay"), None, None)
        for c, v in enumerate(row, start=1):
            if v not in (None, "", 0):
                sheets["workers"].cell(row=r, column=c, value=v)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ══════════════════════════════════════════════════════════════════════════════
# Knowing it again
# ══════════════════════════════════════════════════════════════════════════════

def recognise(table: dict) -> str | None:
    """Which workbook sheet this table is, by its headings. None for any table
    that is not one."""
    cols = {_norm(c) for c in table.get("columns", [])}
    sheet = _norm(table.get("sheet"))
    if sheet == "lists" and "paid by" in cols:
        return "lists"
    if sheet == "start here" and "amount" not in cols:
        return "guide"
    for spec in SHEETS:
        if set(spec["required"]) <= cols:
            return spec["kind"]
    return None


def _actual_mapping(table: dict, spec: dict) -> dict:
    """The spec's columns as the file actually spells them."""
    actual = {_norm(c): c for c in table.get("columns", [])}
    return {field: actual[_norm(col)] for field, col in spec["mapping"].items() if _norm(col) in actual}


def _named_rows(table: dict, column: str | None) -> list:
    if not column:
        return []
    return [r for r in table.get("rows") or []
            if not r.get("_is_total") and str(r.get(column) or "").strip()]


def plan_for(table: dict, kind: str) -> dict:
    """The importer's plan for a recognised table: exact, so no AI is asked."""
    base = {"id": table.get("id"), "confidence": 0.99, "source": "workbook", "workbook": kind,
            "mapping": {}, "event_type": "Expense"}
    if kind == "guide":
        return {**base, "what_it_is": "The instructions page of your AIBOS workbook",
                "import": False, "reason": "Nothing to import on this page."}
    if kind == "lists":
        return {**base, "what_it_is": "The lists behind the workbook's dropdowns",
                "import": False, "reason": "Nothing to import on this page."}
    spec = BY_KIND[kind]
    mapping = _actual_mapping(table, spec)
    if kind in RECORD_KINDS:
        has = bool(_named_rows(table, mapping.get("name")))
    else:
        has = table.get("nonzero_rows", 0) > 0
    return {**base, "what_it_is": spec["what"], "import": has,
            "reason": "" if has else "Nothing filled in on this sheet yet.",
            "event_type": spec["event_type"] or "Expense", "mapping": mapping}


def plan_tables(scanned: dict) -> dict:
    """{table id: plan} for every table that came from the workbook."""
    out = {}
    for t in scanned.get("tables", []):
        kind = recognise(t)
        if kind:
            out[t["id"]] = plan_for(t, kind)
    return out


# ══════════════════════════════════════════════════════════════════════════════
# Filing its lines
# ══════════════════════════════════════════════════════════════════════════════

def _text(v) -> str:
    return " ".join(str(v or "").split())


def _date_key(v) -> str:
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m-%d")
    return _text(v)


def row_key(kind: str, row: dict, mapping: dict, occurrence: int = 1) -> str:
    """What makes this line this line. Two identical sales on one day are two
    lines, so the count of identical lines before it is part of the key: the
    second loaf of bread at K15 on the 3rd is not the first one again."""
    parts = [kind, _date_key(row.get(mapping.get("date")))]
    amt = sheetscan.to_number(row.get(mapping.get("amount")))
    parts.append(f"{amt:.2f}" if amt is not None else "")
    for f in ("counterparty", "description", "quantity", "category"):
        parts.append(_norm(row.get(mapping.get(f))) if mapping.get(f) else "")
    parts.append(str(occurrence))
    return hashlib.sha1("|".join(parts).encode("utf-8")).hexdigest()[:16]


def keys_for(kind: str, rows: list, mapping: dict) -> list:
    """Each row's key, in order, counting repeats so identical lines differ."""
    seen: dict = {}
    out = []
    for r in rows:
        base = row_key(kind, r, mapping, 0)
        seen[base] = seen.get(base, 0) + 1
        out.append(row_key(kind, r, mapping, seen[base]))
    return out


def method_of(v) -> str:
    t = _norm(v)
    if not t:
        return ""
    if "credit" in t:
        return "credit"
    return attach.payment_method_of(t)


def _exact(name: str, records: list) -> dict | None:
    key = _norm(name)
    return next((r for r in records or [] if key and _norm(r.get("name")) == key), None)


def resolve(kind: str, table: dict, mapping: dict, context: dict, known_keys=None,
            heading: str = "") -> dict:
    """Every line of one recognised sheet, as attach.resolve_table returns them:
    {rows, questions, counts}, plus how many lines AIBOS already had.

    A line the owner said everything about is booked as they said it, at the
    confidence of an entry typed on the Record page. Anything that still needs
    working out (a wage to a stranger, an expense with no category, goods not
    on the product list) goes through attach and is asked about, as any import."""
    rows = [r for r in table.get("rows") or []]
    keys = keys_for(kind, rows, mapping)
    known = set(known_keys or ())
    amt_col, desc_col, cp_col = mapping.get("amount"), mapping.get("description"), mapping.get("counterparty")
    qty_col, cat_col = mapping.get("quantity"), mapping.get("category")

    out, already, hand = [], 0, []
    for row, key in zip(rows, keys):
        row = {**row, "_key": key}
        if row.get("_is_total"):
            out.append({**row, "_skip": "total", "_why": "A total line, not a transaction."})
            continue
        amount = sheetscan.to_number(row.get(amt_col)) if amt_col else None
        if not amount:
            out.append({**row, "_skip": "empty", "_why": "No amount on this line."})
            continue
        if key in known:
            already += 1
            out.append({**row, "_skip": "already", "_why": "Already in your books from an earlier upload."})
            continue
        method = method_of(row.get("Paid by"))
        sure = None

        if kind == "sales":
            item, qty = _text(row.get(desc_col)), sheetscan.to_number(row.get(qty_col)) if qty_col else None
            extra = {}
            if _text(row.get(cp_col)):
                extra["customer"] = _text(row.get(cp_col))[:120]
            prod = _exact(item, context.get("products"))
            # Stock moves only with a count: "the day's bread" with no quantity
            # is money, not a number of loaves.
            if item and qty and qty > 0:
                extra["items"] = [prod.get("name") if prod else item[:120]]
                extra["quantities"] = [qty]
                if prod and prod.get("id"):
                    extra["product_id"] = prod["id"]
            sure = {"event_type": "Sale", "payload_extra": extra, "kind": "sale",
                    "match": (prod or {}).get("name"), "confidence": 0.99}

        elif kind == "expenses":
            picked = _CATEGORY.get(_text(row.get(cat_col)).lower()) if cat_col else None
            own = _text(row.get(cat_col)) if cat_col else ""
            if _norm(own) in ("wages", "wage", "salary", "salaries"):
                own = ""           # a wage needs its worker: attach asks who it was
            if not picked and own:
                picked = next((c for c in sum(INDUSTRY_CATEGORIES.values(), ())
                               if c[0].lower() == own.lower()), None)
            if picked and picked[1] == "Loan":
                sure = {"event_type": "Loan", "payload_extra": {"direction": "repayment"},
                        "kind": "service", "match": "loan repayment", "confidence": 0.99}
            elif picked and picked[1] == "TaxPayment":
                sure = {"event_type": "TaxPayment",
                        "payload_extra": {"tax_type": _text(row.get(desc_col))[:60] or "Tax"},
                        "kind": "tax", "match": "tax", "confidence": 0.99}
            elif picked or own:
                sure = {"event_type": "Expense",
                        "payload_extra": {"category": picked[2] if picked else own.lower()[:60]},
                        "kind": "service", "match": picked[0] if picked else own, "confidence": 0.99}

        elif kind == "wages":
            emp = _exact(_text(row.get(cp_col)), context.get("employees"))
            if emp and emp.get("id"):
                sure = {"event_type": "Salary",
                        "payload_extra": {"employee": emp.get("name"), "employee_id": emp["id"],
                                          "category": "salaries"},
                        "kind": "wages", "match": emp.get("name"), "confidence": 0.99}

        elif kind == "stock":
            prod = _exact(_text(row.get(desc_col)), context.get("products"))
            qty = sheetscan.to_number(row.get(qty_col)) if qty_col else None
            if prod and prod.get("id") and qty and qty > 0:
                extra = {"items": [prod.get("name")], "quantities": [qty], "product_id": prod["id"],
                         "category": prod.get("category") or "stock"}
                if _text(row.get(cp_col)):
                    extra["supplier"] = _text(row.get(cp_col))[:120]
                sure = {"event_type": "InventoryReceipt", "payload_extra": extra, "kind": "stock",
                        "match": prod.get("name"), "confidence": 0.99}

        if sure:
            # Nothing on this line was worked out: the importer may book it at
            # the confidence of an entry typed on the Record page.
            sure["sure"] = True
            if method:
                sure["payload_extra"] = {**sure["payload_extra"], "payment_method": method}
            out.append({**row, "_resolved": sure, "_question": None, "_question_key": None})
        else:
            hand.append(len(out))
            out.append(row)

    # The rest go through attach, exactly as any import's rows do.
    questions, counts = [], {"wages": 0, "stock": 0, "service": 0, "tax": 0, "unknown": 0, "sale": 0}
    for r in out:
        v = r.get("_resolved")
        if v:
            counts[v["kind"]] = counts.get(v["kind"], 0) + 1
    if hand:
        ctx = {**context, "default_type": BY_KIND[kind]["event_type"] or "Expense"}
        # The words a question names its line by: the item alone for stock (not
        # "Cement Lafarge", which would become the product's name), the worker
        # alone for wages, and everything said about it for a running cost.
        cols = {"stock": [desc_col], "wages": [cp_col]}.get(kind) or [desc_col, cp_col, "Notes"]
        res = attach.resolve_table([out[i] for i in hand], mapping, ctx,
                                   [c for c in cols if c], heading=heading)
        for i, r in zip(hand, res["rows"]):
            v = r.get("_resolved") or {}
            method = method_of(r.get("Paid by"))
            if method and v:
                v["payload_extra"] = {**(v.get("payload_extra") or {}), "payment_method": method}
            out[i] = r
        questions = res["questions"]
        for k, n in res["counts"].items():
            counts[k] = counts.get(k, 0) + n
    return {"rows": out, "questions": questions, "counts": counts, "already": already}


def heading_for(kind: str) -> str:
    """What the sheet is about, said the way attach reads a table's heading."""
    return {"stock": "stock purchase of goods", "wages": "wages paid to workers",
            "expenses": "", "sales": ""}.get(kind, "")


def records(kind: str, table: dict, mapping: dict, existing: list) -> dict:
    """The Products or Workers sheet as records to add: {"new": [...], "known": [...]}.
    Never a second copy of a name already on file, nor of one twice in the sheet."""
    taken = {_norm(e.get("name")) for e in existing or [] if e.get("name")}
    new, known = [], []
    for r in _named_rows(table, mapping.get("name")):
        data = {}
        for field, col in mapping.items():
            v = r.get(col)
            if v in (None, ""):
                continue
            if field in ("buy_price", "sell_price", "opening_stock", "reorder_level", "basic_pay"):
                n = sheetscan.to_number(v)
                if n is None or n < 0:
                    continue
                data[field] = n
            elif field == "start_date":
                if hasattr(v, "strftime"):
                    data[field] = v.strftime("%Y-%m-%d")
            elif field == "phone":
                # The worker list has no phone column of its own; kept in the notes.
                data["notes"] = f"Phone {_text(v)[:40]}"
            else:
                data[field] = _text(v)[:120]
        name = data.get("name")
        if not name:
            continue
        key = _norm(name)
        if key in taken:
            known.append(name)
            continue
        taken.add(key)
        if kind == "workers":
            data["status"] = "active"
        new.append(data)
    return {"new": new, "known": known}
