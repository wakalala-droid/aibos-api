"""
The AIBOS workbook (starter_book.py): made for the business, known again by its
headings, every line filed as the owner wrote it.

Covered here:
  • the workbook opens in Excel's own reader with every sheet, heading and dropdown
  • the owner's own ways of being paid and biggest costs come first
  • a filled-in workbook is recognised table by table, so no AI is asked
  • a sale, a running cost, a wage and a purchase each land as what they are
  • the Products and Workers sheets go on the lists, never twice
  • the same workbook, added to and uploaded again, brings in only the new lines
  • the routes, end to end
"""

import io
import json
from datetime import datetime

import openpyxl
import pytest

import interview as iv
import sheetscan
import starter_book as sb


def _filled(fill=None, **build_kw) -> bytes:
    """A workbook from build(), with lines typed into it the way an owner would."""
    raw = sb.build(**build_kw)
    wb = openpyxl.load_workbook(io.BytesIO(raw))
    for title, rows in (fill or {}).items():
        ws = wb[title]
        start = ws.max_row + 1 if ws.max_row > 1 and ws.cell(row=ws.max_row, column=1).value else 2
        for i, r in enumerate(rows):
            for j, v in enumerate(r, start=1):
                ws.cell(row=start + i, column=j, value=v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


LINES = {
    "Sales": [
        (datetime(2026, 10, 1), "Banda Builders", "Cement 50kg", 20, 4600, "Mobile money", "Site 4"),
        (datetime(2026, 10, 1), None, "Day's takings", None, 1850, "Cash", None),
        (datetime(2026, 10, 2), None, "Bread", 1, 15, "Cash", None),
        (datetime(2026, 10, 2), None, "Bread", 1, 15, "Cash", None),
        (datetime(2026, 10, 3), "Phiri", "Roofing sheet", 4, 1200, "On credit", None),
    ],
    "Expenses": [
        (datetime(2026, 10, 1), "Landlord", "October rent", "Rent", 6000, "Bank transfer", None),
        (datetime(2026, 10, 4), "Puma", "Diesel for the truck", "Fuel and transport", 900, "Cash", None),
        (datetime(2026, 10, 5), "Access Bank", "Loan instalment", "Loan repayments", 2500, "Bank transfer", None),
        (datetime(2026, 10, 6), "ZESCO", "Electricity units", None, 450, "Mobile money", None),
    ],
    "Stock bought": [
        (datetime(2026, 10, 2), "Lafarge", "Cement 50kg", 100, 19000, "Bank transfer", None),
        (datetime(2026, 10, 3), "Kafue Steel", "Wire nails 4 inch", 10, 800, "Cash", None),
    ],
    "Wages": [
        (datetime(2026, 9, 30), "Mary Banda", 4500, "Mobile money", "September"),
        (datetime(2026, 9, 30), "Chanda Mulenga", 1500, "Cash", "Casual"),
    ],
}

CONTEXT = {
    "employees": [{"id": "e1", "name": "Mary Banda"}],
    "products": [{"id": "p1", "name": "Cement 50kg", "category": "building"},
                 {"id": "p2", "name": "Roofing sheet", "category": "building"}],
    "parties": [],
    "default_type": "Expense",
}


def _table(scanned, sheet):
    return next(t for t in scanned["tables"] if t["sheet"] == sheet)


def _resolved(scanned, sheet, known=None, context=None):
    t = _table(scanned, sheet)
    kind = sb.recognise(t)
    plan = sb.plan_for(t, kind)
    return sb.resolve(kind, t, plan["mapping"], context or CONTEXT, known, sb.heading_for(kind))


# ── The workbook itself ───────────────────────────────────────────────────────

def test_the_workbook_has_a_sheet_for_each_kind_of_money():
    wb = openpyxl.load_workbook(io.BytesIO(sb.build("Mwape Hardware", "Retail")))
    assert wb.sheetnames == ["Start here", "Lists", "Sales", "Expenses", "Stock bought",
                             "Wages", "Products", "Workers"]
    assert wb["Lists"].sheet_state == "hidden"
    assert wb["Start here"]["A1"].value == "Mwape Hardware: AIBOS workbook"
    for spec in sb.SHEETS:
        ws = wb[spec["title"]]
        assert [c.value for c in ws[1]] == [c[0] for c in spec["columns"]]
        assert ws.freeze_panes == "A2"
    assert wb.properties.keywords == sb.TAG


def test_the_dropdowns_are_there_and_never_refuse_the_owner_s_own_words():
    wb = openpyxl.load_workbook(io.BytesIO(sb.build()))
    dvs = wb["Expenses"].data_validations.dataValidation
    by_ref = {str(dv.sqref): dv for dv in dvs}
    assert by_ref["D2:D2000"].formula1.startswith("Lists!$B$2")
    assert all(dv.showErrorMessage is False for dv in dvs)
    sales_refs = {str(dv.sqref) for dv in wb["Sales"].data_validations.dataValidation}
    assert {"C2:C2000", "F2:F2000"} <= sales_refs           # product names, paid by
    wages_refs = {str(dv.sqref) for dv in wb["Wages"].data_validations.dataValidation}
    assert "B2:B2000" in wages_refs                         # worker names


def test_the_owner_s_own_answers_come_first():
    answers = {"paid_by": {"value": ["Mobile money", "Card"]},
               "big_costs": {"value": ["Wages", "Electricity and water", "Rent"]}}
    assert sb.paid_by_for(answers)[:2] == ["Mobile money", "Card"]
    cats = sb.categories_for("Hospitality", answers)
    assert cats[:2] == ["Electricity and water", "Rent"]    # Wages has its own sheet
    assert cats[2:4] == ["Laundry", "Guest supplies"]       # then the industry's own
    assert "Other" in cats and len(cats) == len(set(cats))


def test_the_lists_carry_what_the_business_already_has():
    wb = openpyxl.load_workbook(io.BytesIO(sb.build(
        products=[{"name": "Cement 50kg", "category": "building", "sell_price": 230, "buy_price": 190}],
        workers=[{"name": "Mary Banda", "position": "Cashier", "basic_pay": 4500}])))
    assert [c.value for c in wb["Products"][2]][:5] == ["Cement 50kg", "building", None, 190, 230]
    assert [c.value for c in wb["Workers"][2]][:3] == ["Mary Banda", "Cashier", 4500]


def test_every_word_follows_the_house_style():
    wb = openpyxl.load_workbook(io.BytesIO(sb.build()))
    for row in wb["Start here"].iter_rows(values_only=True):
        for v in row:
            if isinstance(v, str):
                assert "—" not in v and ", and " not in v, v
    for spec in sb.SHEETS:
        assert ", and " not in spec["guide"] and "—" not in spec["what"]


# ── Knowing it again ──────────────────────────────────────────────────────────

def test_a_filled_workbook_is_known_table_by_table():
    scanned = sheetscan.scan(_filled(LINES), "AIBOS workbook.xlsx")
    plans = sb.plan_tables(scanned)
    assert len(plans) == len(scanned["tables"])             # so no AI is asked
    kinds = {sb.recognise(t): t["sheet"] for t in scanned["tables"]}
    assert {"sales", "expenses", "stock", "wages", "guide", "lists"} <= set(kinds)
    sales = plans[_table(scanned, "Sales")["id"]]
    assert sales["import"] and sales["event_type"] == "Sale" and sales["source"] == "workbook"
    assert sales["mapping"] == {"date": "Date", "amount": "Amount", "counterparty": "Customer",
                                "description": "What was sold", "quantity": "Quantity"}
    assert not plans[_table(scanned, "Start here")["id"]]["import"]


def test_empty_sheets_are_left_alone():
    scanned = sheetscan.scan(_filled({"Sales": LINES["Sales"]}), "w.xlsx")
    sheets = {s["name"]: s for s in scanned["sheets"]}
    assert sheets["Expenses"]["skipped"] and sheets["Wages"]["skipped"]


def test_a_sheet_copied_with_its_headings_retyped_is_still_known():
    t = {"sheet": "Sheet1", "columns": ["date", "CUSTOMER", "What  was sold", "Quantity", "amount"]}
    assert sb.recognise(t) == "sales"
    m = sb.plan_for({**t, "id": "x", "nonzero_rows": 1}, "sales")["mapping"]
    assert m["description"] == "What  was sold" and m["amount"] == "amount"


def test_another_spreadsheet_is_not_mistaken_for_the_workbook():
    assert sb.recognise({"sheet": "July", "columns": ["Date", "Description", "Amount"]}) is None
    assert sb.recognise({"sheet": "Lists", "columns": ["Name", "Phone"]}) is None


# ── Filing its lines ──────────────────────────────────────────────────────────

def test_sales_land_as_sales_with_the_stock_they_moved():
    scanned = sheetscan.scan(_filled(LINES), "w.xlsx")
    res = _resolved(scanned, "Sales")
    rows = [r for r in res["rows"] if not r.get("_skip")]
    assert len(rows) == 5 and not res["questions"]
    cement, takings, bread1, bread2, roof = (r["_resolved"] for r in rows)
    assert cement["event_type"] == "Sale" and cement["sure"] and cement["confidence"] == 0.99
    assert cement["payload_extra"] == {"customer": "Banda Builders", "items": ["Cement 50kg"],
                                       "quantities": [20], "product_id": "p1",
                                       "payment_method": "mobile money"}
    # No count, no stock movement: the day's takings is money only.
    assert "items" not in takings["payload_extra"] and takings["payload_extra"]["payment_method"] == "cash"
    assert bread1["payload_extra"]["items"] == ["Bread"] and "product_id" not in bread1["payload_extra"]
    assert roof["payload_extra"]["payment_method"] == "credit"


def test_running_costs_land_under_what_the_owner_picked():
    scanned = sheetscan.scan(_filled(LINES), "w.xlsx")
    rows = [r for r in _resolved(scanned, "Expenses")["rows"] if not r.get("_skip")]
    rent, fuel, loan, zesco = (r["_resolved"] for r in rows)
    assert rent["event_type"] == "Expense" and rent["payload_extra"]["category"] == "rent" and rent["sure"]
    assert rent["payload_extra"]["payment_method"] == "bank"
    assert fuel["payload_extra"]["category"] == "transport"
    assert loan["event_type"] == "Loan" and loan["payload_extra"]["direction"] == "repayment"
    # No category picked: worked out from the words, as any import, and not "sure".
    assert zesco["event_type"] == "Expense" and zesco["payload_extra"]["category"] == "utilities"
    assert not zesco.get("sure") and zesco["payload_extra"]["payment_method"] == "mobile money"


def test_a_wage_to_a_stranger_is_asked_about_by_name():
    scanned = sheetscan.scan(_filled(LINES), "w.xlsx")
    res = _resolved(scanned, "Wages")
    mary, chanda = (r["_resolved"] for r in res["rows"] if not r.get("_skip"))
    assert mary["event_type"] == "Salary" and mary["payload_extra"]["employee_id"] == "e1" and mary["sure"]
    assert chanda["event_type"] == "Salary" and not chanda.get("sure")
    assert [(q["type"], q["name"]) for q in res["questions"]] == [("unknown_worker", "Chanda Mulenga")]


def test_stock_lands_on_the_product_and_an_unknown_item_is_asked_about_by_its_own_name():
    scanned = sheetscan.scan(_filled(LINES), "w.xlsx")
    res = _resolved(scanned, "Stock bought")
    cement, nails = (r["_resolved"] for r in res["rows"] if not r.get("_skip"))
    assert cement["event_type"] == "InventoryReceipt" and cement["sure"]
    assert cement["payload_extra"]["quantities"] == [100] and cement["payload_extra"]["supplier"] == "Lafarge"
    q = res["questions"][0]
    assert q["type"] == "unknown_product" and q["name"] == "Wire nails 4 inch"


def test_a_worker_on_the_workbook_s_own_list_is_not_a_stranger():
    ctx = {**CONTEXT, "employees": CONTEXT["employees"] + [{"id": None, "name": "Chanda Mulenga"}]}
    scanned = sheetscan.scan(_filled(LINES), "w.xlsx")
    assert _resolved(scanned, "Wages", context=ctx)["questions"] == []


def test_the_same_workbook_again_brings_in_only_the_new_lines():
    first = sheetscan.scan(_filled(LINES), "w.xlsx")
    res = _resolved(first, "Sales")
    keys = [r["_key"] for r in res["rows"] if not r.get("_skip")]
    assert len(set(keys)) == 5                     # the two loaves are two lines

    more = {**LINES, "Sales": LINES["Sales"] + [
        (datetime(2026, 10, 2), None, "Bread", 1, 15, "Cash", None),      # a third loaf
        (datetime(2026, 10, 9), "Banda Builders", "Cement 50kg", 10, 2300, "Cash", None)]}
    again = _resolved(sheetscan.scan(_filled(more), "w.xlsx"), "Sales", known=keys)
    fresh = [r for r in again["rows"] if not r.get("_skip")]
    assert again["already"] == 5 and len(fresh) == 2
    assert {r["What was sold"] for r in fresh} == {"Bread", "Cement 50kg"}


def test_the_lists_go_on_file_once():
    fill = {"Products": [("Cement 50kg", "building", "bag", 190, 230, 40, 10),
                         ("Wire nails 4 inch", "hardware", "kg", 70, 95, None, None),
                         ("wire nails 4 inch", None, None, None, 95, None, None)],
            "Workers": [("Chanda Mulenga", "Shop hand", 1500, "0977 123456", datetime(2025, 3, 1))]}
    scanned = sheetscan.scan(_filled(fill), "w.xlsx")
    p = _table(scanned, "Products")
    got = sb.records("products", p, sb.plan_for(p, "products")["mapping"], CONTEXT["products"])
    assert got["known"] == ["Cement 50kg", "wire nails 4 inch"]     # on file, and listed twice
    assert got["new"] == [{"name": "Wire nails 4 inch", "category": "hardware", "unit": "kg",
                           "buy_price": 70, "sell_price": 95}]
    w = _table(scanned, "Workers")
    got = sb.records("workers", w, sb.plan_for(w, "workers")["mapping"], CONTEXT["employees"])
    assert got["new"] == [{"name": "Chanda Mulenga", "position": "Shop hand", "basic_pay": 1500,
                           "notes": "Phone 0977 123456", "start_date": "2025-03-01",
                           "status": "active"}]


# ── The routes ────────────────────────────────────────────────────────────────

def _client(db, path):
    from test_document_import import _client as client
    return client(db, path)


def _books():
    from test_books_integrity import _fresh
    db = _fresh()
    db.rows["businesses"].append({"id": "b1", "owner_id": "u1", "name": "Mwape Hardware",
                                  "industry": "Retail", "currency": "ZMW", "is_default": True})
    db.rows["employees"] = [{"id": "e1", "user_id": "u1", "name": "Mary Banda", "status": "active"}]
    db.rows["products"] = [
        {"id": "p1", "user_id": "u1", "business_id": "b1", "name": "Cement 50kg", "category": "building"},
        {"id": "p2", "user_id": "u1", "business_id": "b1", "name": "Roofing sheet", "category": "building"}]
    return db


@pytest.fixture
def restore():
    import main
    real = main.get_db
    yield
    main.get_db = real
    main.app.dependency_overrides.clear()


def test_the_workbook_downloads_made_for_the_business(restore):
    db = _books()
    db.rows["business_interviews"] = [{"id": "i1", "user_id": "u1", "business_id": "b1",
                                       "answers": {"paid_by": {"value": ["Mobile money"]}}}]
    client = _client(db, "/templates/workbook.xlsx")
    r = client.get("/templates/workbook.xlsx")
    assert r.status_code == 200, r.text
    assert 'filename="Mwape Hardware AIBOS workbook.xlsx"' in r.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb["Lists"]["A2"].value == "Mobile money"
    assert wb["Products"]["A2"].value == "Cement 50kg" and wb["Workers"]["A2"].value == "Mary Banda"


def test_the_scan_files_the_workbook_without_the_ai(restore, monkeypatch):
    import doc_ai
    monkeypatch.setattr(doc_ai, "plan", lambda *a, **k: pytest.fail("the AI was asked"))
    fill = {**LINES, "Workers": [("Chanda Mulenga", "Shop hand", 1500, None, None)]}
    client = _client(_books(), "/documents/scan")
    r = client.post("/documents/scan", files={"file": ("w.xlsx", _filled(fill), "x")})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["workbook"] is True and body["ai"] is False
    by_sheet = {t["sheet"]: t for t in body["tables"]}
    assert by_sheet["Sales"]["workbook"] == "sales" and by_sheet["Sales"]["import"]
    assert by_sheet["Workers"]["records"] == {"new": ["Chanda Mulenga"], "new_count": 1, "known_count": 0}
    # Chanda is on the workbook's own Workers sheet: not asked about.
    assert not any(q["type"] == "unknown_worker" for q in body["questions"])
    assert any(q["type"] == "unknown_product" and q["name"] == "Wire nails 4 inch"
               for q in body["questions"])


def test_the_import_books_the_workbook_and_a_second_upload_adds_only_new_lines(restore):
    db = _books()
    fill = {**LINES, "Workers": [("Chanda Mulenga", "Shop hand", 1500, None, None)],
            "Products": [("Wire nails 4 inch", "hardware", "kg", 70, 95, None, None)]}
    client = _client(db, "/documents/import")
    r = client.post("/documents/import", files={"file": ("w.xlsx", _filled(fill), "x")},
                    data={"tables": "[]", "answers": "{}", "currency": "ZMW"})
    assert r.status_code == 200, r.text
    body = r.json()
    assert sorted(body["recorded"]["employees"]) == ["Chanda Mulenga"]
    assert body["recorded"]["products"] == ["Wire nails 4 inch"]
    events = db.rows["business_events"]
    assert len(events) == 13
    # What the owner said in full is booked like an entry on the Record page.
    sale = next(e for e in events if e["event_type"] == "Sale" and e["payload"]["amount"] == 4600)
    assert sale["status"] == "confirmed" and sale["payload"]["product_id"] == "p1"
    chanda = next(e for e in events if e["event_type"] == "Salary" and e["payload"]["amount"] == 1500)
    assert chanda["payload"]["employee_id"] and chanda["status"] == "confirmed"
    nails = next(e for e in events if e["payload"].get("amount") == 800)
    assert nails["event_type"] == "InventoryReceipt" and nails["status"] == "confirmed"
    zesco = next(e for e in events if e["payload"].get("amount") == 450)
    assert zesco["status"] == "pending"              # its category was worked out, so it waits

    more = {**fill, "Sales": LINES["Sales"] + [
        (datetime(2026, 10, 9), "Banda Builders", "Cement 50kg", 10, 2300, "Cash", None)]}
    r = client.post("/documents/import", files={"file": ("w.xlsx", _filled(more), "x")},
                    data={"tables": "[]", "answers": "{}", "currency": "ZMW"})
    assert r.status_code == 200, r.text
    assert r.json()["saved_count"] == 1 and r.json()["already"] == 13
    assert len(db.rows["business_events"]) == 14
    assert sum(1 for e in db.rows["employees"] if e["name"] == "Chanda Mulenga") == 1


def test_a_workbook_with_nothing_new_says_so(restore):
    db = _books()
    client = _client(db, "/documents/import")
    data = {"tables": "[]", "answers": "{}", "currency": "ZMW"}
    first = _filled({"Sales": LINES["Sales"][:1]})
    assert client.post("/documents/import", files={"file": ("w.xlsx", first, "x")}, data=data).status_code == 200
    # Saved again by Excel, so the bytes differ but the lines do not.
    wb = openpyxl.load_workbook(io.BytesIO(first))
    wb.properties.title = "saved again"
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post("/documents/import", files={"file": ("w.xlsx", buf.getvalue(), "x")}, data=data)
    assert r.status_code == 400 and "already in your books" in r.json()["detail"]
